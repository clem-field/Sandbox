import os
import openpyxl
import pandas as pd

def extract_info(intro_sheet):
    system_owner = None
    authors = None
    reference = None
    smes = None
    rows = list(intro_sheet.iter_rows(values_only=True))
    for i, row in enumerate(rows):
        if row and row[0]:
            cell_str = str(row[0]).lower()
            if "system owner of the author's area" in cell_str:
                if i + 1 < len(rows) and rows[i + 1][0]:
                    system_owner = rows[i + 1][0]
            elif "assigned to this checklist" in cell_str:
                if i + 1 < len(rows) and rows[i + 1][0]:
                    authors = str(rows[i + 1][0]).replace('\n', ', ')
            elif "source reference document" in cell_str:
                if i + 1 < len(rows) and rows[i + 1][0]:
                    reference = rows[i + 1][0]
            elif "subject matter experts" in cell_str:
                if i + 1 < len(rows) and rows[i + 1][0]:
                    smes = str(rows[i + 1][0]).replace('\n', ', ')
    return system_owner, authors, reference, smes

directory = '.'  # Change to your directory path if needed
data = []

for filename in os.listdir(directory):
    if filename.endswith('.xlsx') and not filename.endswith('_clean.xlsx'):
        path = os.path.join(directory, filename)
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            
            # Copy Checklist sheet to new workbook, removing second row if needed
            if 'Checklist' in wb.sheetnames:
                clean_wb = openpyxl.Workbook()
                clean_sheet = clean_wb.active
                clean_sheet.title = 'Checklist'
                source_sheet = wb['Checklist']
                rows = list(source_sheet.iter_rows(values_only=True))
                
                # Check if second row's control ID contains the specified text
                rows_to_copy = rows
                if len(rows) >= 2 and rows[1][0] and "This number uniquely identifies each control" in str(rows[1][0]):
                    rows_to_copy = rows[:1] + rows[2:]  # Skip second row
                
                # Copy rows to clean sheet
                for row_idx, row in enumerate(rows_to_copy, start=1):
                    for col_idx, value in enumerate(row, start=1):
                        clean_sheet.cell(row=row_idx, column=col_idx).value = value
                
                clean_filename = filename.rsplit('.xlsx', 1)[0] + '_clean.xlsx'
                clean_wb.save(os.path.join(directory, clean_filename))
            
            # Extract from Introduction
            if 'Introduction' in wb.sheetnames:
                intro_sheet = wb['Introduction']
                system_owner, authors, reference, smes = extract_info(intro_sheet)
                data.append({
                    'file_name': filename,
                    'system_owner': system_owner,
                    'authors': authors,
                    'references': reference,
                    'smes': smes
                })
        except Exception as e:
            print(f"Error processing {filename}: {e}")

# Aggregate into table and save
if data:
    df = pd.DataFrame(data)
    df.to_excel(os.path.join(directory, 'summary.xlsx'), index=False)
    print("Summary table saved to summary.xlsx")
else:
    print("No files processed.")