import argparse
from docx import Document
import pandas as pd
import sys

def normalize_text(text):
    """Normalize text by converting to lowercase and removing extra whitespace."""
    return ' '.join(str(text).lower().split()) if text else ''

def parse_docx_table(doc_path):
    try:
        print(f"Attempting to open document: {doc_path}")
        doc = Document(doc_path)
    except Exception as e:
        print(f"Error opening document {doc_path}: {e}")
        return

    print(f"Found {len(doc.tables)} table(s) in the document")
    expected_headers = ['Item #', 'Configuration Item', 'Action / Recommended Parameters', 'Rational/Remediation', 'Windows', 'Unix', 'Level & score']
    norm_expected_headers = [normalize_text(h) for h in expected_headers]
    all_data = []
    matching_tables = 0

    for table_idx, table in enumerate(doc.tables):
        # Get headers
        headers = [cell.text.strip() for cell in table.rows[0].cells if cell.text.strip()]
        norm_headers = [normalize_text(h) for h in headers]
        print(f"Table {table_idx} headers (raw): {headers}")
        print(f"Table {table_idx} headers (normalized): {norm_headers}")

        # Check if headers match (case-insensitive, ignoring extra whitespace)
        if norm_headers[:len(norm_expected_headers)] == norm_expected_headers:
            print(f"Table {table_idx} matches expected headers")
            matching_tables += 1
            table_data = []
            for row_idx, row in enumerate(table.rows[1:], start=1):
                cells = row.cells
                if len(cells) < 7:
                    print(f"Skipping row {row_idx} in table {table_idx}: insufficient columns ({len(cells)})")
                    continue
                item_num = cells[0].text.strip()
                if not item_num:  # Skip empty rows
                    print(f"Skipping row {row_idx} in table {table_idx}: empty Item #")
                    continue
                config_item = cells[1].text.strip()
                action = cells[2].text.strip()
                rat_rem = cells[3].text.strip()
                windows = cells[4].text.strip()
                unix = cells[5].text.strip()
                level_score = cells[6].text.strip()
                
                # Split rat_rem into rationale, remediation, audit
                rationale = ''
                remediation = ''
                audit = ''
                parts = rat_rem.split('\n')
                current = None
                for line in parts:
                    line = line.strip()
                    if not line:
                        continue
                    lower_line = normalize_text(line)
                    if lower_line.startswith('rationale:') or lower_line.startswith('rational:'):
                        current = 'rationale'
                        rationale = line.split(':', 1)[1].strip() if ':' in line else ''
                    elif lower_line.startswith('remediation:'):
                        current = 'remediation'
                        remediation = line.split(':', 1)[1].strip() if ':' in line else ''
                    elif lower_line.startswith('audit:'):
                        current = 'audit'
                        audit = line.split(':', 1)[1].strip() if ':' in line else ''
                    elif current:
                        if current == 'rationale':
                            rationale += ' ' + line
                        elif current == 'remediation':
                            remediation += ' ' + line
                        elif current == 'audit':
                            audit += ' ' + line
                
                table_data.append({
                    'Item #': item_num,
                    'Configuration Item': config_item,
                    'Action / Recommended Parameters': action,
                    'Rationale': rationale.strip(),
                    'Remediation': remediation.strip(),
                    'Audit': audit.strip(),
                    'Windows': windows,
                    'Unix': unix,
                    'Level / Score': level_score
                })
            
            print(f"Table {table_idx} processed {len(table_data)} valid rows")
            all_data.extend(table_data)
        else:
            print(f"Table {table_idx} does not match expected headers")

    if all_data:
        try:
            df = pd.DataFrame(all_data)
            output_path = doc_path.replace('.docx', '_parsed.xlsx')
            df.to_excel(output_path, index=False, engine='openpyxl')
            print(f"Spreadsheet saved to {output_path} with {len(all_data)} rows from {matching_tables} table(s)")
        except Exception as e:
            print(f"Error saving Excel file: {e}")
    else:
        print(f"No valid data found in {matching_tables} matching table(s). Check table contents and structure.")

def main():
    parser = argparse.ArgumentParser(description="Parse a CIS DOCX file and convert specified tables to Excel.")
    parser.add_argument('-i', '--input', required=True, help="Path to the input DOCX file")
    args = parser.parse_args()
    
    if not args.input.endswith('.docx'):
        print("Error: Input file must be a .docx file")
        return
    
    try:
        import docx
        import pandas
        print(f"Using python-docx version: {docx.__version__}")
        print(f"Using pandas version: {pandas.__version__}")
    except AttributeError as e:
        print(f"Error accessing package versions: {e}")
    
    parse_docx_table(args.input)

if __name__ == "__main__":
    main()